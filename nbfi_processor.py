# -*- coding: utf-8 -*-
"""
NBFI Returns PDF → Excel কনভার্টার
======================================
Karmasangsthan Bank-এর "FORM NBFI RETURNS" PDF রিপোর্ট (Account-wise Loan/Lease/Advances
তথ্য, multi-page টেবিল) থেকে ডেটা বের করে ব্যাংকের নিজস্ব NBFI রিপোর্টিং টেমপ্লেটের
(৩৪-কলাম ফরম্যাট) মতো একটা .xlsx ফাইল বানায়।

PDF-এর ১-২৬ নম্বর কলাম (DATED থেকে OVERDUE_AMOUNT) সরাসরি কপি হয়। বাকি কলামগুলো:
    27. Sanc/Disb               = SANCTION_AMOUNT - DISBURSED_AMOUNT           (ফর্মুলা)
    28. Overdue/Outstanding     = OUTSTANDING_AMOUNT - OVERDUE_AMOUNT          (ফর্মুলা)
    29. Project_No              = ইউজারের দেওয়া মান (আপাতত ডিফল্ট 23)
    30. Unic_Kormosuchi_Code_NO = Account Number থেকে বের করা (নিচে দেখুন)
    31. Br_Name                 = ইউজারের ইনপুট (PDF-এ নেই)
    32. RM_Office_Name          = ইউজারের ইনপুট (PDF-এ নেই)
    33. Out Form                = OPENING_BALANCE + DISBURSED_AMOUNT + ACCRUED_INTEREST
                                  + OTHER_CHARGES - RECOVERED_AMOUNT - ADJUSTMENT_AMOUNT
                                  - WRITE_OFF_AMOUNT                            (ফর্মুলা)
    34. Out Diff                = OUTSTANDING_AMOUNT - Out Form                 (ফর্মুলা)

Unic_Kormosuchi_Code_NO হিসাবের নিয়ম (hint অনুযায়ী):
    Account Number-কে string বানিয়ে (দশমিক থাকলে দশমিক ও তার পরের অংশ বাদ দিয়ে),
    শেষ চার ডিজিটের ঠিক আগের দুই ডিজিট নেওয়া হয়। যেমন 1120308010342 -> শেষ চার
    "0342", তার আগের দুই "01" -> 1। এটা মূল Excel টেমপ্লেটের ১০৫০+ real সারির বিপরীতে
    যাচাই করে দেখা হয়েছে (~৯৯.৯% মেলে)।
"""
import re
from datetime import datetime

import pdfplumber
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

HEADERS = [
    "DATED", "FI_ID", "FI_BRANCH_ID", "ACCOUNT_NUMBER", "ACCOUNT_HOLDER'S_NAME",
    "DATE_OF_BIRTH", "GENDER_CODE", "UNIQUE_ID_TYPE", "UNIQUE_ID",
    "ECONOMIC_SECTOR_CODE", "ECONOMIC_PURPOSE_CODE", "INDUSTRY_SCALE_CODE",
    "Security/COLLATERAL_CODE", "PRODUCT_TYPE_CODE", "LOAN_CLASS_CODE",
    "INTEREST_RATE", "SANCTION_AMOUNT", "OPENING_BALANCE", "DISBURSED_AMOUNT",
    "RECOVERED_AMOUNT", "ACCRUED_INTEREST", "OTHER_CHARGES", "ADJUSTMENT_AMOUNT",
    "WRITE_OFF_AMOUNT", "OUTSTANDING_AMOUNT", "OVERDUE_AMOUNT",
    "Sanc/Disb", "Overdue/\nOutstanding", "Project_No (Category of Loan)",
    "Unic_Kormosuchi_Code_NO", "Br_Name", "RM_Office_Name", "Out Form", "Out Diff",
]

_AMOUNT_COLS = {17, 18, 19, 20, 21, 22, 23, 24, 25, 26, 27, 28, 33, 34}  # 1-based


def _num(v):
    try:
        return float(str(v).replace(",", "").strip())
    except (TypeError, ValueError):
        return 0.0


def _as_int_or_raw(v):
    s = str(v).strip() if v is not None else ""
    return int(s) if s.lstrip("-").isdigit() else v


def _parse_ddmmyy(s):
    """'31-08-26' -> datetime(2026, 8, 31)"""
    s = (s or "").strip()
    try:
        return datetime.strptime(s, "%d-%m-%y")
    except ValueError:
        return None


def _parse_dob(s):
    """'31-MAY-84' -> datetime(1984, 5, 31)"""
    s = (s or "").strip()
    try:
        return datetime.strptime(s, "%d-%b-%y")
    except ValueError:
        return None


def compute_kormosuchi_code(account_number):
    """Unic_Kormosuchi_Code_NO -- Account Number-এর (দশমিক বাদ দিয়ে) শেষ চার
    ডিজিটের ঠিক আগের দুই ডিজিট থেকে বের করা হয়। যেমন 1120308010342 -> শেষ চার
    '0342', তার আগের দুই '01' -> 1। দশমিক থাকলে (যেমন 1120308010598.2) সেই
    দশমিক অংশ বাদ দিয়ে হিসাব করা হয়। Account Number ছোট/অস্বাভাবিক হলে None।"""
    if account_number is None:
        return None
    s = str(account_number).strip()
    if "." in s:
        s = s.split(".")[0]
    s = re.sub(r"\D", "", s)
    if len(s) < 6:
        return None
    return int(s[-6:-4])


def parse_nbfi_header(path):
    """PDF-এর প্রথম পেজ থেকে ব্যাংকের নাম, শাখার নাম ও রিপোর্টিং পিরিয়ড বের করে।"""
    with pdfplumber.open(path) as pdf:
        txt = pdf.pages[0].extract_text() or ""
    m_bank = re.search(r"^(.*?)\s+FORM NBFI RETURNS", txt)
    m_branch = re.search(r"\n([^\n]*Branch)\n", txt)
    m_period = re.search(r"From\s+(\d{2}/\d{2}/\d{4})\s+to\s+(\d{2}/\d{2}/\d{4})", txt)
    return {
        "bank_name": m_bank.group(1).strip() if m_bank else "",
        "branch_name": m_branch.group(1).strip() if m_branch else "",
        "period_from": m_period.group(1) if m_period else "",
        "period_to": m_period.group(2) if m_period else "",
    }


def parse_nbfi_returns_pdf(path):
    """PDF-এর প্রতি পেজ থেকে টেবিল বের করে row dict-এর লিস্ট রিটার্ন করে। প্রতিটা
    dict-এ HEADERS-এর প্রথম ২৬টা key (DATED...OVERDUE_AMOUNT) থাকে -- মান PDF থেকে
    যথাযথভাবে parse করে (সংখ্যা/তারিখ হিসেবে) বসানো হয়। হেডার/নাম্বারিং রো বাদ যায়।"""
    rows = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            for tbl in page.find_tables():
                for r in tbl.extract():
                    if not r or r[0] is None:
                        continue
                    sl = str(r[0]).strip()
                    if not sl.isdigit():
                        continue  # হেডার সারি / কলাম-নাম্বারিং সারি বাদ
                    if len(r) < 27:
                        continue
                    acc_raw = (r[4] or "").strip()
                    if "." in acc_raw:
                        try:
                            acc_num = float(acc_raw)
                        except ValueError:
                            acc_num = acc_raw
                    else:
                        acc_num = _as_int_or_raw(acc_raw)
                    name = re.sub(r"\s+", " ", (r[5] or "").strip())
                    row = {
                        "DATED": _parse_ddmmyy(r[1]),
                        "FI_ID": _as_int_or_raw(r[2]),
                        "FI_BRANCH_ID": _as_int_or_raw(r[3]),
                        "ACCOUNT_NUMBER": acc_num,
                        "ACCOUNT_HOLDER'S_NAME": name,
                        "DATE_OF_BIRTH": _parse_dob(r[6]),
                        "GENDER_CODE": _as_int_or_raw(r[7]),
                        "UNIQUE_ID_TYPE": _as_int_or_raw(r[8]),
                        "UNIQUE_ID": (r[9] or "").strip(),
                        "ECONOMIC_SECTOR_CODE": _as_int_or_raw(r[10]),
                        "ECONOMIC_PURPOSE_CODE": _as_int_or_raw(r[11]),
                        "INDUSTRY_SCALE_CODE": _as_int_or_raw(r[12]),
                        "Security/COLLATERAL_CODE": _as_int_or_raw(r[13]),
                        "PRODUCT_TYPE_CODE": _as_int_or_raw(r[14]),
                        "LOAN_CLASS_CODE": _as_int_or_raw(r[15]),
                        "INTEREST_RATE": _num(r[16]),
                        "SANCTION_AMOUNT": _num(r[17]),
                        "OPENING_BALANCE": _num(r[18]),
                        "DISBURSED_AMOUNT": _num(r[19]),
                        "RECOVERED_AMOUNT": _num(r[20]),
                        "ACCRUED_INTEREST": _num(r[21]),
                        "OTHER_CHARGES": _num(r[22]),
                        "ADJUSTMENT_AMOUNT": _num(r[23]),
                        "WRITE_OFF_AMOUNT": _num(r[24]),
                        "OUTSTANDING_AMOUNT": _num(r[25]),
                        "OVERDUE_AMOUNT": _num(r[26]),
                    }
                    rows.append(row)
    return rows


def build_nbfi_excel(rows, out_path, project_no=23, br_name="", rm_office_name="",
                      title_text=None):
    """rows: parse_nbfi_returns_pdf()-এর রেজাল্ট। ব্যাংকের ৩৪-কলাম NBFI রিপোর্টিং
    টেমপ্লেটের মতো header (row 1-4) + data (row 5 থেকে) সহ .xlsx বানায়।"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "NBFI Report"

    green_fill = PatternFill("solid", fgColor="00B050")
    center_wrap = Alignment(horizontal="center", vertical="center", wrap_text=True)
    bold = Font(bold=True)

    if title_text is None:
        title_text = "Account-wise details information of Loan/Lease and Advances"
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(HEADERS))
    title_cell = ws.cell(row=1, column=1, value=title_text)
    title_cell.font = Font(bold=True, size=12)
    title_cell.alignment = Alignment(horizontal="center")

    for idx, h in enumerate(HEADERS, start=1):
        cell = ws.cell(row=2, column=idx, value=h)
        cell.font = bold
        cell.alignment = center_wrap
        cell.fill = green_fill
        ws.merge_cells(start_row=2, start_column=idx, end_row=3, end_column=idx)
        num_cell = ws.cell(row=4, column=idx, value=idx)
        num_cell.font = bold
        num_cell.alignment = Alignment(horizontal="center")

    start_row = 5
    for i, row in enumerate(rows):
        r = start_row + i
        for idx, key in enumerate(HEADERS[:26], start=1):
            ws.cell(row=r, column=idx, value=row.get(key))
        ws.cell(row=r, column=27, value=f"=Q{r}-S{r}")
        ws.cell(row=r, column=28, value=f"=Y{r}-Z{r}")
        ws.cell(row=r, column=29, value=project_no)
        ws.cell(row=r, column=30, value=compute_kormosuchi_code(row.get("ACCOUNT_NUMBER")))
        ws.cell(row=r, column=31, value=br_name)
        ws.cell(row=r, column=32, value=rm_office_name)
        ws.cell(row=r, column=33, value=f"=R{r}+S{r}+U{r}+V{r}-T{r}-W{r}-X{r}")
        ws.cell(row=r, column=34, value=f"=Y{r}-AG{r}")

        ws.cell(row=r, column=1).number_format = "DD-MMM-YY"
        ws.cell(row=r, column=6).number_format = "DD-MMM-YY"
        for col in _AMOUNT_COLS:
            ws.cell(row=r, column=col).number_format = "#,##0"

    widths = [10, 7, 12, 16, 26, 12, 8, 8, 20, 9, 9, 9, 9, 9, 9, 9, 11, 11, 11, 11,
              10, 9, 10, 10, 12, 11, 9, 9, 20, 11, 16, 14, 11, 9]
    for idx, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = w

    ws.freeze_panes = "A5"

    # --- যোগফল সারি -- আসল ব্যাংক টেমপ্লেটে যেভাবে আছে ঠিক সেভাবেই: Q(17) থেকে AH(34)
    # পর্যন্ত প্রতিটা amount কলামের যোগফল (Project_No আর Unic_Kormosuchi_Code_NO বাদে,
    # যেহেতু ওগুলো কোড, টাকার অংক না) ---
    last_row = start_row + len(rows) - 1
    total_row = last_row + 1 if rows else start_row

    label_cell = ws.cell(row=total_row, column=1, value="Total")
    label_cell.font = bold
    sum_cols = [c for c in range(17, len(HEADERS) + 1) if c not in (29, 30)]
    for c in sum_cols:
        col_letter = get_column_letter(c)
        cell = ws.cell(row=total_row, column=c,
                        value=f"=SUM({col_letter}{start_row}:{col_letter}{last_row})")
        cell.font = bold
        cell.number_format = "#,##0"

    # --- বাড়তি সারাংশ টেবিল: Product Type Code (কলাম N)-অনুযায়ী সংখ্যা এবং সেই
    # রো-গুলোর Overdue/Outstanding (কলাম AB)-এর যোগফল, যেমন: 21031 | 60 | 4,500,000 ---
    n_col = get_column_letter(14)   # N -- PRODUCT_TYPE_CODE
    ab_col = get_column_letter(28)  # AB -- Overdue/Outstanding

    def _criteria(code):
        return str(code) if isinstance(code, (int, float)) else f'"{code}"'

    codes = sorted(
        {row.get("PRODUCT_TYPE_CODE") for row in rows if row.get("PRODUCT_TYPE_CODE") is not None},
        key=lambda v: (isinstance(v, str), v),
    )

    summary_start = total_row + 3
    heading_cell = ws.cell(row=summary_start, column=1,
                            value="Product Type Code (Column N) অনুযায়ী সারাংশ")
    heading_cell.font = Font(bold=True, size=11)

    hdr_row = summary_start + 1
    for idx, h in enumerate(["Product Type Code", "সংখ্যা (Count)", "Overdue/Outstanding (AB) যোগফল"], start=1):
        c = ws.cell(row=hdr_row, column=idx, value=h)
        c.font = bold
        c.fill = green_fill
        c.alignment = center_wrap

    for i, code in enumerate(codes):
        r = hdr_row + 1 + i
        ws.cell(row=r, column=1, value=code)
        ws.cell(row=r, column=2,
                value=f"=COUNTIF({n_col}{start_row}:{n_col}{last_row},{_criteria(code)})")
        sum_cell = ws.cell(
            row=r, column=3,
            value=f"=SUMIF({n_col}{start_row}:{n_col}{last_row},{_criteria(code)},"
                  f"{ab_col}{start_row}:{ab_col}{last_row})",
        )
        sum_cell.number_format = "#,##0"

    wb.save(out_path)
    return out_path
